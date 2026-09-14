"""Group 4 — spreadsheets & tabular data: the model gets the MAP (sheets,
shapes, columns, types, samples), never a raw dump of a million rows. Tiny
files are the exception — those are worth inlining whole.

xlsx/xlsm → openpyxl (read-only mode)   xls → pandas+xlrd   ods → pandas+odfpy
csv/tsv → pandas                        parquet/orc → pyarrow
"""

from typing import Any

from ..detect import ext_of
from ..result import HandlerResult

# A CSV at or under this many rows is dumped whole instead of sampled.
FULL_DUMP_ROWS = 100
SAMPLE_ROWS = 5
MAX_COLS_LISTED = 60


def _fmt_cell(v: Any) -> str:
    s = "" if v is None else str(v)
    s = s.replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= 60 else s[:57] + "…"


def _table(headers: list[str], rows: list[list[Any]]) -> str:
    head = "| " + " | ".join(_fmt_cell(h) for h in headers) + " |"
    sep = "|" + "---|" * len(headers)
    body = "\n".join("| " + " | ".join(_fmt_cell(c) for c in r) + " |" for r in rows)
    return f"{head}\n{sep}\n{body}" if body else f"{head}\n{sep}"


def _df_section(name: str, df, total_rows: int | None = None) -> tuple[str, dict]:
    """Describe one sheet/table from a pandas DataFrame."""
    rows = total_rows if total_rows is not None else len(df)
    cols = list(df.columns.astype(str))
    dtypes = [str(t) for t in df.dtypes]
    listed = cols[:MAX_COLS_LISTED]
    col_lines = "\n".join(
        f"- `{c}` ({t})" for c, t in zip(listed, dtypes[: len(listed)])
    )
    if len(cols) > MAX_COLS_LISTED:
        col_lines += f"\n- … and {len(cols) - MAX_COLS_LISTED} more columns"

    sample = df.head(SAMPLE_ROWS)
    section = (
        f"## {name} — {rows:,} rows × {len(cols)} columns\n\n"
        f"Columns:\n{col_lines}\n\n"
        f"Sample rows:\n{_table(cols[:MAX_COLS_LISTED], sample.values.tolist())}"
    )
    meta = {"name": name, "rows": rows, "columns": len(cols)}
    return section, meta


def _csv(path: str, filename: str, sep: str) -> HandlerResult:
    import pandas as pd

    # Count data rows cheaply, then decide dump vs schema.
    with open(path, "rb") as f:
        total = sum(1 for _ in f)
    data_rows = max(total - 1, 0)

    if data_rows <= FULL_DUMP_ROWS:
        df = pd.read_csv(path, sep=sep)
        content = (
            f"# {filename} — full contents ({len(df):,} rows)\n\n"
            + _table(list(df.columns.astype(str)), df.values.tolist())
        )
        return HandlerResult(
            status="ready",
            processor_group="spreadsheet",
            content=content,
            meta={"rows": len(df), "columns": len(df.columns), "fullDump": True},
        )

    df = pd.read_csv(path, sep=sep, nrows=200)  # enough to infer types
    section, sheet_meta = _df_section(filename, df, total_rows=data_rows)
    content = (
        f"# {filename} — schema overview\n\n{section}\n\n"
        "_Schema + sample only; run code against the stored file for full data._"
    )
    return HandlerResult(
        status="ready",
        processor_group="spreadsheet",
        content=content,
        meta={"rows": data_rows, "columns": sheet_meta["columns"], "fullDump": False},
    )


def _excel_like(path: str, filename: str, engine: str | None) -> HandlerResult:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None, nrows=200, engine=engine)
    sections: list[str] = []
    sheet_metas: list[dict] = []
    for name, df in sheets.items():
        # nrows caps the sample read; get true counts for xlsx via openpyxl.
        section, m = _df_section(str(name), df)
        sections.append(section)
        sheet_metas.append(m)

    # For OOXML we can read exact dimensions cheaply in read-only mode.
    if engine is None or engine == "openpyxl":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            dims = {ws.title: ws.max_row for ws in wb.worksheets}
            wb.close()
            for m in sheet_metas:
                if m["name"] in dims and dims[m["name"]]:
                    m["rows"] = max(dims[m["name"]] - 1, 0)
        except Exception:  # noqa: BLE001 — dimensions stay sample-based
            pass

    content = (
        f"# {filename} — workbook overview ({len(sections)} sheet"
        f"{'s' if len(sections) != 1 else ''})\n\n" + "\n\n".join(sections) + "\n\n"
        "_Schema + samples only; run code against the stored file for full data._"
    )
    return HandlerResult(
        status="ready",
        processor_group="spreadsheet",
        content=content,
        meta={"sheets": sheet_metas},
    )


def _columnar(path: str, filename: str, ext: str) -> HandlerResult:
    import pyarrow.parquet as pq

    if ext == "parquet":
        pf = pq.ParquetFile(path)
        schema = pf.schema_arrow
        rows = pf.metadata.num_rows
        sample = pf.read_row_group(0).to_pandas().head(SAMPLE_ROWS) if pf.metadata.num_row_groups else None
    else:  # orc
        from pyarrow import orc

        of = orc.ORCFile(path)
        schema = of.schema
        rows = of.nrows
        sample = of.read().to_pandas().head(SAMPLE_ROWS)

    cols = [f"- `{f.name}` ({f.type})" for f in schema]
    content = (
        f"# {filename} — columnar file schema\n\n"
        f"{rows:,} rows × {len(cols)} columns\n\nColumns:\n" + "\n".join(cols)
    )
    if sample is not None:
        content += "\n\nSample rows:\n" + _table(
            list(sample.columns.astype(str)), sample.values.tolist()
        )
    return HandlerResult(
        status="ready",
        processor_group="spreadsheet",
        content=content,
        meta={"rows": rows, "columns": len(cols)},
    )


def handle(path: str, filename: str, job: object) -> HandlerResult:
    ext = ext_of(filename)
    try:
        if ext == "csv":
            return _csv(path, filename, ",")
        if ext == "tsv":
            return _csv(path, filename, "\t")
        if ext in ("xlsx", "xlsm"):
            return _excel_like(path, filename, "openpyxl")
        if ext == "xls":
            return _excel_like(path, filename, "xlrd")
        if ext == "ods":
            return _excel_like(path, filename, "odf")
        if ext in ("parquet", "orc"):
            return _columnar(path, filename, ext)
        # avro / .numbers — no reader wired; metadata with an honest note.
        return HandlerResult(
            status="unsupported",
            processor_group="spreadsheet",
            meta={"note": f"No reader for .{ext}; stored and addressable by path."},
        )
    except Exception as e:  # noqa: BLE001 — corrupt/exotic file
        return HandlerResult(
            status="failed",
            processor_group="spreadsheet",
            error=f"{type(e).__name__}: {e}",
        )
